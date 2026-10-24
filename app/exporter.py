import io
import json
from typing import Any, Dict, Tuple

from docx import Document
from docx.shared import Pt, RGBColor

from app.audio_utils import format_timestamp
from app.grounding import readable_summaries


def _hex_to_rgb(hex_color: str) -> Tuple[int, int, int]:
    h = hex_color.lstrip("#")
    if len(h) == 6:
        return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return 79, 70, 229


# Headings of exported documents, in the language of the document: the translation for a translated
# export, otherwise the language of the recording. Other languages use English.
_LABELS: Dict[str, Dict[str, str]] = {
    "ru": {
        "untitled": "Запись", "date": "Дата", "processed": "Дата обработки", "duration": "Длительность",
        "model": "Модель", "participants": "Участники", "unknown": "не определены", "overview": "Обзор",
        "transcript": "Расшифровка", "translation": "Перевод",
        "mt_note": "Машинный перевод ({lang}), модель {model}", "mt_stale": "устарело: {stale}, без перевода: {missing}",
    },
    "en": {
        "untitled": "Recording", "date": "Date", "processed": "Processed", "duration": "Duration",
        "model": "Model", "participants": "Participants", "unknown": "not identified", "overview": "Overview",
        "transcript": "Transcript", "translation": "Translation",
        "mt_note": "Machine translation ({lang}), model {model}", "mt_stale": "outdated: {stale}, untranslated: {missing}",
    },
    "de": {
        "untitled": "Aufnahme", "date": "Datum", "processed": "Verarbeitet", "duration": "Dauer",
        "model": "Modell", "participants": "Teilnehmer", "unknown": "nicht erkannt", "overview": "Überblick",
        "transcript": "Transkript", "translation": "Übersetzung",
        "mt_note": "Maschinelle Übersetzung ({lang}), Modell {model}", "mt_stale": "veraltet: {stale}, nicht übersetzt: {missing}",
    },
    "fr": {
        "untitled": "Enregistrement", "date": "Date", "processed": "Traité le", "duration": "Durée",
        "model": "Modèle", "participants": "Participants", "unknown": "non identifiés", "overview": "Synthèse",
        "transcript": "Transcription", "translation": "Traduction",
        "mt_note": "Traduction automatique ({lang}), modèle {model}", "mt_stale": "obsolètes : {stale}, non traduits : {missing}",
    },
    "es": {
        "untitled": "Grabación", "date": "Fecha", "processed": "Procesado", "duration": "Duración",
        "model": "Modelo", "participants": "Participantes", "unknown": "no identificados", "overview": "Resumen",
        "transcript": "Transcripción", "translation": "Traducción",
        "mt_note": "Traducción automática ({lang}), modelo {model}", "mt_stale": "desactualizados: {stale}, sin traducir: {missing}",
    },
}


def _labels(meeting: Dict[str, Any]) -> Dict[str, str]:
    lang = (meeting.get("translation_lang") or meeting.get("language") or "ru").lower()
    return _LABELS.get(lang, _LABELS["en"])


def _translation_note(meeting: Dict[str, Any], L: Dict[str, str]) -> str:
    info = meeting.get("translation_info")
    if not info:
        return meeting.get("translation_note") or ""
    note = L["mt_note"].format(lang=info["lang"].upper(), model=info["model"])
    if info.get("stale") or info.get("missing"):
        note += "; " + L["mt_stale"].format(stale=info.get("stale", 0), missing=info.get("missing", 0))
    return note


def _shown(meeting: Dict[str, Any]) -> tuple:
    """Whether the export shows the time stamps and the names (export_build.with_options)."""
    opts = meeting.get("_export") or {}
    return opts.get("times", True), opts.get("names", True)


def export_markdown(meeting: Dict[str, Any]) -> str:
    times, names = _shown(meeting)
    L = _labels(meeting)
    title = meeting.get("title") or L["untitled"]
    created_at = meeting.get("created_at", "")
    duration_str = format_timestamp(meeting.get("duration", 0.0))
    speakers = meeting.get("speakers") or {}
    segments = meeting.get("segments") or []
    summaries = readable_summaries(meeting)

    lines = [
        f"# {title}",
        "",
        f"- **{L['processed']}:** {created_at}",
        f"- **{L['duration']}:** {duration_str}",
        f"- **{L['model']}:** `{meeting.get('model_name', 'large-v3-turbo')}`",
        f"- **{L['participants']} ({len(speakers)}):** "
        + (", ".join(s.get("name", k) for k, s in speakers.items()) or L["unknown"]),
    ]
    note = _translation_note(meeting, L)
    if note:
        lines.append(f"- **{L['translation']}:** {note}")
    lines.append("")

    if summaries.get("protocol"):
        lines.extend(
            [
                "---",
                "",
                summaries["protocol"],
                "",
            ]
        )
    elif summaries.get("tldr"):
        lines.extend(
            [
                "---",
                "",
                summaries["tldr"],
                "",
            ]
        )

    lines.extend(
        [
            "---",
            "",
            f"## {L['transcript']}",
            "",
        ]
    )

    for seg in segments:
        spk_id = seg.get("speaker", "SPEAKER_00")
        spk_name = speakers.get(spk_id, {}).get("name", spk_id)
        ts_start = format_timestamp(seg.get("start", 0.0))
        ts_end = format_timestamp(seg.get("end", 0.0))
        head = " ".join(x for x in (f"[{ts_start} – {ts_end}]" if times else "", spk_name if names else "") if x)
        if head:
            lines.append(f"**{head}:**")
        lines.append(f"{seg.get('text', '').strip()}")
        if seg.get("secondary_text"):
            lines.append(f"*{seg['secondary_text'].strip()}*")
        lines.append("")

    return "\n".join(lines)


def export_txt(meeting: Dict[str, Any]) -> str:
    times, names = _shown(meeting)
    L = _labels(meeting)
    title = meeting.get("title") or L["untitled"]
    created_at = meeting.get("created_at", "")
    duration_str = format_timestamp(meeting.get("duration", 0.0))
    speakers = meeting.get("speakers") or {}
    segments = meeting.get("segments") or []
    summaries = readable_summaries(meeting)

    lines = [
        title,
        f"{L['date']}: {created_at} | {L['duration']}: {duration_str}",
    ]
    note = _translation_note(meeting, L)
    if note:
        lines.append(f"{L['translation']}: {note}")
    lines += ["=" * 60, ""]

    if summaries.get("protocol"):
        lines.extend(
            [
                f"{L['overview'].upper()}:",
                "-" * 40,
                summaries["protocol"],
                "",
                "=" * 60,
                "",
            ]
        )

    lines.append(f"{L['transcript'].upper()}:")
    lines.append("-" * 40)

    for seg in segments:
        spk_id = seg.get("speaker", "SPEAKER_00")
        spk_name = speakers.get(spk_id, {}).get("name", spk_id)
        ts_start = format_timestamp(seg.get("start", 0.0))
        head = " ".join(x for x in (f"[{ts_start}]" if times else "", f"{spk_name}:" if names else "") if x)
        lines.append(f"{head} {seg.get('text', '').strip()}".strip())
        if seg.get("secondary_text"):
            lines.append(f"    {seg['secondary_text'].strip()}")

    return "\n".join(lines)


def export_json(meeting: Dict[str, Any]) -> str:
    payload = {
        "id": meeting.get("id"),
        "title": meeting.get("title"),
        "filename": meeting.get("filename"),
        "duration": meeting.get("duration"),
        "language": meeting.get("language"),
        "model_name": meeting.get("model_name"),
        "created_at": meeting.get("created_at"),
        "speakers": meeting.get("speakers", {}),
        "summaries": meeting.get("summaries", {}),
        "segments": meeting.get("segments", []),
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


def export_docx(meeting: Dict[str, Any]) -> bytes:
    times, names = _shown(meeting)
    doc = Document()

    L = _labels(meeting)
    title = meeting.get("title") or L["untitled"]
    speakers = meeting.get("speakers") or {}
    segments = meeting.get("segments") or []
    summaries = readable_summaries(meeting)

    doc.add_heading(title, level=0)

    meta_p = doc.add_paragraph()
    meta_p.add_run(f"{L['date']}: {meeting.get('created_at', '')}  |  ").italic = True
    meta_p.add_run(
        f"{L['duration']}: {format_timestamp(meeting.get('duration', 0.0))}  |  "
    ).italic = True
    participant_names = ", ".join(s.get("name", k) for k, s in speakers.items())
    meta_p.add_run(f"{L['participants']}: {participant_names or L['unknown']}").italic = True
    note = _translation_note(meeting, L)
    if note:
        meta_p.add_run(f"\n{L['translation']}: {note}").italic = True

    if summaries.get("protocol"):
        doc.add_heading(L["overview"], level=1)
        for line in summaries["protocol"].splitlines():
            clean = line.strip()
            if not clean:
                continue
            if clean.startswith("## "):
                doc.add_heading(clean.replace("## ", ""), level=2)
            elif clean.startswith("### "):
                doc.add_heading(clean.replace("### ", ""), level=3)
            else:
                doc.add_paragraph(clean)

    doc.add_heading(L["transcript"], level=1)

    for seg in segments:
        spk_id = seg.get("speaker", "SPEAKER_00")
        spk_info = speakers.get(spk_id, {})
        spk_name = spk_info.get("name", spk_id)
        spk_color = spk_info.get("color", "#4f46e5")
        r, g, b = _hex_to_rgb(spk_color)

        ts_start = format_timestamp(seg.get("start", 0.0))
        p = doc.add_paragraph()

        if times:
            ts_run = p.add_run(f"[{ts_start}] ")
            ts_run.font.size = Pt(9.5)
            ts_run.font.color.rgb = RGBColor(120, 120, 130)

        if names:
            spk_run = p.add_run(f"{spk_name}: ")
            spk_run.bold = True
            spk_run.font.color.rgb = RGBColor(r, g, b)

        txt_run = p.add_run(seg.get("text", "").strip())
        txt_run.font.size = Pt(11)

        if seg.get("secondary_text"):
            tr_p = doc.add_paragraph()
            tr_p.paragraph_format.left_indent = Pt(18)
            tr_run = tr_p.add_run(seg["secondary_text"].strip())
            tr_run.italic = True
            tr_run.font.size = Pt(10.5)
            tr_run.font.color.rgb = RGBColor(90, 90, 110)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
